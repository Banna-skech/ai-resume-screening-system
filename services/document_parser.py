"""文档解析服务 —— PDF (pdfminer) + DOCX (python-docx) 文本提取"""

import re
import logging
from pathlib import Path

logger = logging.getLogger(__name__)


def clean_text(text: str) -> str:
    """清洗提取的文本：规范化空白、去除页码、去除多余空行

    Args:
        text: 原始文本

    Returns:
        清洗后的文本
    """
    # 统一换行符
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # 去除多余空行（3 个以上换行 → 2 个）
    text = re.sub(r"\n{3,}", "\n\n", text)

    # 去除独立的页码行（纯数字行）
    text = re.sub(r"\n\s*\d{1,3}\s*\n", "\n", text)

    # 去除行首行尾空白
    lines = [line.strip() for line in text.splitlines()]
    text = "\n".join(lines)

    # 去除首尾空白
    text = text.strip()

    return text


def parse_pdf(file_path: Path) -> str:
    """从 PDF 文件中提取文本

    使用 pdfminer.six 进行文本层提取。
    对于扫描件 PDF（文本层为空），返回空字符串供上层判断是否需要 OCR。

    Args:
        file_path: PDF 文件路径

    Returns:
        提取并清洗后的文本
    """
    try:
        from pdfminer.high_level import extract_text
    except ImportError:
        raise ImportError("请安装 pdfminer.six: pip install pdfminer.six")

    text = extract_text(str(file_path))
    if not text or len(text.strip()) < 30:
        logger.warning(f"PDF 文本层内容极少（{len(text.strip())} 字符），可能是扫描件: {file_path.name}")
        return ""

    return clean_text(text)


def parse_docx(file_path: Path) -> str:
    """从 DOCX 文件中提取文本

    Args:
        file_path: DOCX 文件路径

    Returns:
        提取并清洗后的文本
    """
    try:
        from docx import Document
    except ImportError:
        raise ImportError("请安装 python-docx: pip install python-docx")

    doc = Document(str(file_path))

    paragraphs = []
    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text.strip())

    # 也提取表格中的文本
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())

    text = "\n".join(paragraphs)
    return clean_text(text)


def parse_resume(file_path: Path) -> str:
    """统一的简历文件解析入口

    根据文件扩展名自动选择解析器。

    Args:
        file_path: 简历文件路径

    Returns:
        提取并清洗后的文本

    Raises:
        ValueError: 不支持的文件格式
    """
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        text = parse_pdf(file_path)
        if not text:
            logger.warning(
                f"⚠️ {file_path.name} 可能是扫描件 PDF，文本层内容为空。"
                f"如需 OCR 支持，请安装 pdf2image + pytesseract。"
            )
        return text

    elif suffix in (".docx", ".doc"):
        return parse_docx(file_path)

    else:
        raise ValueError(f"不支持的文件格式: {suffix}，支持的格式: .pdf, .docx")


def try_ocr_pdf(file_path: Path, languages: str = "chi_sim+eng") -> str:
    """对扫描件 PDF 进行 OCR 识别（可选功能）

    需要额外安装: pip install pdf2image pytesseract
    以及系统安装: poppler-utils, tesseract-ocr

    Args:
        file_path: PDF 文件路径
        languages: Tesseract 语言设置

    Returns:
        OCR 识别后的文本
    """
    try:
        from pdf2image import convert_from_path
        import pytesseract
    except ImportError:
        raise ImportError(
            "OCR 功能需要额外依赖: pip install pdf2image pytesseract\n"
            "以及系统安装 poppler-utils 和 tesseract-ocr（含中文语言包）"
        )

    images = convert_from_path(str(file_path), dpi=300)
    texts = []
    for i, img in enumerate(images):
        page_text = pytesseract.image_to_string(img, lang=languages)
        texts.append(page_text)

    full_text = "\n".join(texts)
    return clean_text(full_text)
