"""Excel 报告生成服务 —— 多 Sheet 报告：汇总 + 明细 + 技能矩阵"""

import logging
from pathlib import Path
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from models.match_result import MatchResult
from config.settings import settings
from config.constants import (
    EXCEL_FILENAME_PREFIX,
    EXCEL_SHEET_SUMMARY,
    EXCEL_SHEET_DETAIL,
    EXCEL_SHEET_SKILLS,
    DECISION_PASS,
    DECISION_REVIEW,
    DECISION_REJECT,
    DECISION_LABELS,
    DIMENSION_LABELS,
)

logger = logging.getLogger(__name__)

# 样式定义
HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

PASS_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")  # 绿色
REVIEW_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")  # 黄色
REJECT_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")  # 红色

CELL_FONT = Font(name="微软雅黑", size=10)
CELL_ALIGNMENT = Alignment(vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

DECISION_FILL_MAP = {
    DECISION_PASS: PASS_FILL,
    DECISION_REVIEW: REVIEW_FILL,
    DECISION_REJECT: REJECT_FILL,
}


class ReportGenerator:
    """Excel 报告生成器

    生成包含以下 Sheet 的工作簿：
    - 汇总表：所有候选人的概况
    - 详细评分：每项评分明细
    - 技能矩阵：候选人 × 技能 交叉表
    """

    def __init__(self, report_dir: str | None = None):
        self.report_dir = Path(report_dir or settings.report_dir)
        self.report_dir.mkdir(parents=True, exist_ok=True)

    def generate(self, results: list[MatchResult], job_title: str = "") -> Path:
        """生成 Excel 报告

        Args:
            results: 匹配结果列表
            job_title: 岗位名称（用于文件名）

        Returns:
            生成的 Excel 文件路径
        """
        wb = Workbook()

        # Sheet 1: 汇总表
        self._build_summary_sheet(wb, results)

        # Sheet 2: 详细评分
        self._build_detail_sheet(wb, results)

        # Sheet 3: 技能矩阵
        self._build_skills_matrix(wb, results)

        # 删除默认创建的空白 sheet
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

        # 保存
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        title_part = f"_{job_title}" if job_title else ""
        filename = f"{EXCEL_FILENAME_PREFIX}{title_part}_{timestamp}.xlsx"
        filepath = self.report_dir / filename
        wb.save(str(filepath))

        logger.info(f"Excel 报告已生成: {filepath}")
        return filepath

    def _build_summary_sheet(self, wb: Workbook, results: list[MatchResult]):
        """构建汇总表"""
        ws = wb.active
        ws.title = EXCEL_SHEET_SUMMARY

        headers = [
            "序号", "候选人姓名", "文件名", "综合评分",
            "技能匹配", "经验匹配", "背景匹配", "意向匹配",
            "决策", "决策说明", "标签",
        ]
        self._write_header(ws, headers)

        for i, result in enumerate(results):
            row = i + 2
            ws.cell(row=row, column=1, value=i + 1)
            ws.cell(row=row, column=2, value=result.candidate_name)
            ws.cell(row=row, column=3, value=result.resume_file)
            ws.cell(row=row, column=4, value=result.overall_score)
            ws.cell(row=row, column=5, value=self._get_dim_score(result, "skills"))
            ws.cell(row=row, column=6, value=self._get_dim_score(result, "experience"))
            ws.cell(row=row, column=7, value=self._get_dim_score(result, "background"))
            ws.cell(row=row, column=8, value=self._get_dim_score(result, "intention"))
            ws.cell(row=row, column=9, value=DECISION_LABELS.get(result.decision, result.decision))
            ws.cell(row=row, column=10, value=result.decision_reason)

            # 标签汇总
            tag_texts = []
            for tag in result.tags:
                icon = "✓" if tag.match else "✗"
                tag_texts.append(f"{icon}{tag.tag}")
            ws.cell(row=row, column=11, value=", ".join(tag_texts))

            # 行样式
            fill = DECISION_FILL_MAP.get(result.decision)
            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=row, column=col)
                cell.font = CELL_FONT
                cell.alignment = CELL_ALIGNMENT
                cell.border = THIN_BORDER
                if fill:
                    cell.fill = fill

        # 列宽
        widths = [6, 15, 25, 10, 10, 10, 10, 10, 10, 35, 40]
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(i + 1)].width = w

        # 冻结首行
        ws.freeze_panes = "A2"

        # 自动筛选
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(results) + 1}"

    def _build_detail_sheet(self, wb: Workbook, results: list[MatchResult]):
        """构建详细评分表"""
        ws = wb.create_sheet(EXCEL_SHEET_DETAIL)

        headers = [
            "候选人姓名", "维度", "分数", "权重",
            "评分项", "影响", "评语",
        ]
        self._write_header(ws, headers)

        row = 2
        for result in results:
            for ds in result.dimension_scores:
                dim_label = DIMENSION_LABELS.get(ds.dimension, ds.dimension)
                if ds.details:
                    for detail in ds.details:
                        ws.cell(row=row, column=1, value=result.candidate_name)
                        ws.cell(row=row, column=2, value=dim_label)
                        ws.cell(row=row, column=3, value=ds.score)
                        ws.cell(row=row, column=4, value=f"{ds.weight:.0%}")
                        ws.cell(row=row, column=5, value=detail.reason)
                        impact_cn = {"positive": "✅ 正面", "negative": "❌ 负面", "neutral": "➖ 中性"}.get(
                            detail.impact, detail.impact
                        )
                        ws.cell(row=row, column=6, value=impact_cn)
                        ws.cell(row=row, column=7, value="")

                        for col in range(1, len(headers) + 1):
                            cell = ws.cell(row=row, column=col)
                            cell.font = CELL_FONT
                            cell.border = THIN_BORDER
                        row += 1
                else:
                    # 该维度没有明细，仍然输出一行
                    ws.cell(row=row, column=1, value=result.candidate_name)
                    ws.cell(row=row, column=2, value=dim_label)
                    ws.cell(row=row, column=3, value=ds.score)
                    ws.cell(row=row, column=4, value=f"{ds.weight:.0%}")
                    ws.cell(row=row, column=5, value="（无明细）")
                    ws.cell(row=row, column=6, value="")
                    ws.cell(row=row, column=7, value="")
                    for col in range(1, len(headers) + 1):
                        cell = ws.cell(row=row, column=col)
                        cell.font = CELL_FONT
                        cell.border = THIN_BORDER
                    row += 1

        widths = [15, 12, 8, 8, 40, 10, 40]
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(i + 1)].width = w

        ws.freeze_panes = "A2"

    def _build_skills_matrix(self, wb: Workbook, results: list[MatchResult]):
        """构建技能矩阵（候选人 × 技能标签）"""
        ws = wb.create_sheet(EXCEL_SHEET_SKILLS)

        # 收集所有技能标签
        all_skill_tags: list[str] = []
        for result in results:
            for tag in result.tags:
                if tag.category in ("skill", "education"):
                    if tag.tag not in all_skill_tags:
                        all_skill_tags.append(tag.tag)

        # 写表头
        ws.cell(row=1, column=1, value="候选人")
        for j, tag in enumerate(all_skill_tags):
            ws.cell(row=1, column=j + 2, value=tag)

        # 写匹配矩阵
        for i, result in enumerate(results):
            row = i + 2
            ws.cell(row=row, column=1, value=result.candidate_name)

            candidate_tags = {t.tag: t.match for t in result.tags}
            for j, tag in enumerate(all_skill_tags):
                match_val = candidate_tags.get(tag)
                cell = ws.cell(row=row, column=j + 2)
                if match_val is True:
                    cell.value = "✓"
                    cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                elif match_val is False:
                    cell.value = "✗"
                    cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                else:
                    cell.value = ""
                cell.font = CELL_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = THIN_BORDER

            # 首列样式
            ws.cell(row=row, column=1).font = CELL_FONT
            ws.cell(row=row, column=1).border = THIN_BORDER

        # 列宽
        ws.column_dimensions["A"].width = 15
        for j in range(len(all_skill_tags)):
            ws.column_dimensions[get_column_letter(j + 2)].width = 12

        ws.freeze_panes = "B2"

    def _write_header(self, ws, headers: list[str]):
        """写入表头行"""
        for j, header in enumerate(headers):
            cell = ws.cell(row=1, column=j + 1, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGNMENT
            cell.border = THIN_BORDER

    def _get_dim_score(self, result: MatchResult, dimension: str) -> float:
        """获取指定维度的分数"""
        for ds in result.dimension_scores:
            if ds.dimension == dimension:
                return ds.score
        return 0.0
